import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from amh_assignment import AmhAssignment


class AmhWorkbookTests(unittest.TestCase):
    def make_assignment(self, source, **options):
        assignment = AmhAssignment.__new__(AmhAssignment)
        assignment.config = {"files": {"business_output": {
            "path": str(source), "sheet_name": "Cases", **options}}}
        assignment.as_of_date = pd.Timestamp("2026-10-26")
        assignment.cm_df = pd.DataFrame([
            {"cm_id": "001", "cm_name": "Alice & Lee", "team": "TH A"},
            {"cm_id": "002", "cm_name": "Bob", "team": "TH B"},
        ])
        assignment.warnings = []
        return assignment

    def results(self):
        return pd.DataFrame([
            {"Task ID": "400", "Assigned Analyst ID 1": "002", "Assigned Analyst ID 2": "001"},
            {"Task ID": "300", "Assigned Analyst ID 1": "001", "Assigned Analyst ID 2": "002"},
        ])

    def test_records_and_transposed_preserve_workbook_and_match_customer(self):
        for transposed in (False, True):
            for suffix in (".xlsx", ".xlsm"):
                with self.subTest(transposed=transposed, suffix=suffix), tempfile.TemporaryDirectory() as directory:
                    source = Path(directory) / f"cases{suffix}"
                    book = Workbook()
                    sheet = book.active
                    sheet.title = "Cases"
                    rows = [
                        ["Customer_ID_S1", "CIN", "Assign by TH", "CM Name", "Case_Manager_Staff_ID_S1"],
                        ["C1", "000300", "old TH", "old CM", ""],
                        ["C2", "000400", "old TH", "old CM", ""],
                        ["C3", "000500", "unchanged TH", "unchanged CM", ""],
                        ["C4", "000300", "old TH", "old CM", ""],
                    ]
                    for row in zip(*rows) if transposed else rows:
                        sheet.append(list(row))
                    target = "B4" if transposed else "D2"
                    sheet[target].fill = PatternFill("solid", fgColor="00FF00")
                    book.create_sheet("Other")["A1"] = "=1+2"
                    book.save(source)
                    book.close()
                    # Sentinel stands in for VBA and other opaque workbook parts.
                    with ZipFile(source, "a") as archive:
                        archive.writestr("xl/vbaProject.bin", b"opaque VBA content")
                    original = source.read_bytes()
                    assignment = self.make_assignment(source, orientation="auto")
                    output = assignment.export_assignment_workbook(self.results())
                    self.assertEqual(output.name, f"cases_assigned_2026-10-26{suffix}")
                    self.assertEqual(source.read_bytes(), original)
                    with ZipFile(source) as before, ZipFile(output) as after:
                        self.assertEqual(before.namelist(), after.namelist())
                        for name in before.namelist():
                            if name != "xl/worksheets/sheet1.xml":
                                self.assertEqual(before.read(name), after.read(name), name)
                    filled = load_workbook(output)
                    sheet = filled["Cases"]
                    addresses = ["B3", "B4", "C3", "C4", "D4", "E4"] if transposed else ["C2", "D2", "C3", "D3", "D4", "D5"]
                    self.assertEqual([sheet[a].value for a in addresses],
                                     ["TH A", "Alice & Lee", "TH B", "Bob", "unchanged CM", "Alice & Lee"])
                    self.assertEqual(sheet[target].fill.fgColor.rgb, "0000FF00")
                    filled.close()
                    if transposed:
                        assignment.raw_case_df = pd.read_excel(source, sheet_name="Cases")
                        assignment.prepare_case_list()
                        self.assertEqual(assignment.case_df["case_id"].tolist(), ["C1", "C2", "C3", "C4"])

    def test_source_cannot_be_overwritten(self):
        assignment = self.make_assignment(Path("input.xlsx"))
        assignment.config["filled_output_file"] = "input.xlsx"
        with self.assertRaisesRegex(ValueError, "cannot overwrite"):
            assignment.export_assignment_workbook(self.results())

    def test_missing_fields_fail_without_producing_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "input.xlsx"
            book = Workbook()
            book.active.title = "Cases"
            book.active.append(["CIN", "CM Name"])
            book.active.append(["300", "old"])
            book.save(source)
            book.close()
            assignment = self.make_assignment(source)
            with self.assertRaisesRegex(ValueError, "Assign by TH"):
                assignment.export_assignment_workbook(self.results())
            self.assertEqual(list(Path(directory).iterdir()), [source])


if __name__ == "__main__":
    unittest.main()
